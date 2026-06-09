import os
from typing import Union, Tuple, Dict

import openpyxl
import numpy as np


class ReadOnlyExcel:

    def __init__(self, file_name:str):
        self.workbook:openpyxl.Workbook = openpyxl.load_workbook(file_name, read_only=True, data_only=True)
        self.file_name = os.path.abspath(file_name).replace("\\", "/")
        self.base_name = os.path.basename(file_name)

        self.active_sheet = self.workbook.active
        self.active_sheet_index:int = 0
        self.sheet_names = self.workbook.sheetnames
        self.active_sheet_name:str = self.sheet_names[0]
        self.data_map:Dict[str, np.ndarray] = {}

    def sheet(self, index:Union[int,str])->None:
        if isinstance(index, int):
            self.active_sheet = self.workbook[self.sheet_names[index]]
            self.active_sheet_name = self.sheet_names[index]
        elif isinstance(index, str):
            self.active_sheet = self.workbook[index]
            self.active_sheet_name = index

    @property
    def sheet_name(self)->str:
        return self.active_sheet.title
    
    @sheet_name.setter
    def sheet_name(self, name:str)->None:
        self.active_sheet.title = name

    def __getitem__(self, index:Tuple[int, int]):
        row = index[0]
        column = index[1]
        return self.active_sheet.cell(row=row+1, column=column+1).value

    def is_empty(self, row:int, col:int)->bool:
        value = self[row, col].value
        return value is None or (isinstance(value, str) and value == '')

    @property
    def rows(self)->int:
        return self.active_sheet.max_row
    
    @property
    def cols(self)->int:
        return self.active_sheet.max_column
    
    def full_data(self, stage=None)->np.ndarray:
        if self.active_sheet_name in self.data_map:
            if stage is not None:
                stage.progress = 1
            return self.data_map[self.active_sheet_name]

        data = []
        actual_max_row = min(self.active_sheet.max_row, 100000)
        actual_max_col = min(self.active_sheet.max_column, 10000)

        i:int = 0
        for row in self.active_sheet.iter_rows(
            min_row=1,
            max_row=actual_max_row,
            min_col=1,
            max_col=actual_max_col,
            values_only=True
        ):
            if all(cell is None for cell in row):
                break

            data.append(row)
            i += 1
            if stage is not None:
                stage.progress = i / actual_max_row

        self.data_map[self.active_sheet_name] = np.array(data)
        return self.data_map[self.active_sheet_name]